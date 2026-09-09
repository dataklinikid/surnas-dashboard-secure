from __future__ import annotations

import hashlib
import re
import unicodedata
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import IntegrityError, transaction
from openpyxl import load_workbook

from aggregate.models import SurveyAccess, SurveyPSU, SurveyPSUFrame


MAX_FRAME_BYTES = 5 * 1024 * 1024
HEADER_ALIASES = {
    "NO": "psu_number",
    "DESA/KELURAHAN": "village",
    "KECAMATAN": "district",
    "KABUPATEN/KOTA": "regency",
    "DAPIL DPRRI": "dpr_ri_constituency",
    "DAPIL DPR RI": "dpr_ri_constituency",
    "PROVINSI": "province",
    "STATUS": "urban_rural",
    "RESPONDEN": "target_n",
    "TARGET RESPONDEN": "target_n",
    "TARGET N": "target_n",
    "NO KUES": "questionnaire_range",
    "NO KUESIONER": "questionnaire_range",
}
REQUIRED_FIELDS = frozenset(HEADER_ALIASES.values())


class PSUFrameError(ValueError):
    pass


@dataclass(frozen=True)
class PSUFrameCandidate:
    source_name: str
    file_sha256: str
    rows: tuple[dict, ...]
    report: dict


def _header(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().upper())


def normalize_location(*parts) -> str:
    normalized = []
    for part in parts:
        value = unicodedata.normalize("NFKD", str(part or ""))
        value = "".join(char for char in value if not unicodedata.combining(char))
        normalized.append(re.sub(r"[^A-Z0-9]+", " ", value.upper()).strip())
    return "|".join(normalized)


def _positive_integer(value, label: str, excel_row: int) -> int:
    try:
        decimal_value = Decimal(str(value).strip())
    except (InvalidOperation, TypeError, ValueError):
        raise PSUFrameError(f"Baris {excel_row}: {label} harus berupa bilangan bulat.")
    if decimal_value != decimal_value.to_integral_value():
        raise PSUFrameError(f"Baris {excel_row}: {label} harus berupa bilangan bulat.")
    number = int(decimal_value)
    if number <= 0:
        raise PSUFrameError(f"Baris {excel_row}: {label} harus lebih besar dari nol.")
    return number


def _questionnaire_range(value, excel_row: int) -> tuple[int, int]:
    text = str(value or "").strip()
    numbers = re.findall(r"\d+", text)
    if len(numbers) == 1:
        start = end = int(numbers[0])
    elif len(numbers) == 2:
        start, end = map(int, numbers)
    else:
        raise PSUFrameError(
            f"Baris {excel_row}: NO KUES harus berupa satu nomor atau rentang, misalnya 1--10."
        )
    if start <= 0 or end < start:
        raise PSUFrameError(f"Baris {excel_row}: rentang NO KUES tidak valid.")
    return start, end


def inspect_psu_frame(uploaded_file) -> PSUFrameCandidate:
    raw = uploaded_file.read()
    try:
        uploaded_file.seek(0)
    except (AttributeError, OSError):
        pass
    source_name = str(getattr(uploaded_file, "name", "psu_frame.xlsx"))
    if not source_name.lower().endswith(".xlsx"):
        raise PSUFrameError("Frame PSU harus berupa file .xlsx.")
    if not raw:
        raise PSUFrameError("File frame PSU kosong.")
    if len(raw) > MAX_FRAME_BYTES:
        raise PSUFrameError("File frame PSU melebihi batas 5 MB.")
    try:
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise PSUFrameError("File XLSX tidak dapat dibaca atau rusak.") from exc
    try:
        if "PSU" not in workbook.sheetnames:
            raise PSUFrameError("Sheet bernama PSU tidak ditemukan.")
        sheet = workbook["PSU"]
        header_row = None
        columns = {}
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 20), values_only=True),
            start=1,
        ):
            candidate = {
                HEADER_ALIASES[label]: index
                for index, value in enumerate(values)
                if (label := _header(value)) in HEADER_ALIASES
            }
            if REQUIRED_FIELDS.issubset(candidate):
                header_row = row_number
                columns = candidate
                break
        if header_row is None:
            raise PSUFrameError(
                "Header frame tidak lengkap. Wajib: NO, DESA/KELURAHAN, KECAMATAN, "
                "KABUPATEN/KOTA, DAPIL DPRRI, PROVINSI, STATUS, RESPONDEN, dan NO KUES."
            )

        parsed = []
        seen_psu = set()
        seen_locations = set()
        intervals = []
        warnings = []
        for excel_row, values in enumerate(
            sheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            first = values[columns["psu_number"]] if columns["psu_number"] < len(values) else None
            if first is None or str(first).strip() == "":
                continue
            if _header(first) == "TOTAL":
                break
            psu_number = _positive_integer(first, "NO", excel_row)
            if psu_number in seen_psu:
                raise PSUFrameError(f"Baris {excel_row}: nomor PSU {psu_number} duplikat.")
            seen_psu.add(psu_number)

            def value(field):
                index = columns[field]
                return str(values[index] if index < len(values) and values[index] is not None else "").strip()

            text_fields = {
                field: value(field)
                for field in (
                    "village", "district", "regency", "dpr_ri_constituency",
                    "province", "urban_rural",
                )
            }
            missing = [field for field, item in text_fields.items() if not item]
            if missing:
                raise PSUFrameError(
                    f"Baris {excel_row}: field wajib kosong ({', '.join(missing)})."
                )
            target_n = _positive_integer(
                values[columns["target_n"]], "RESPONDEN", excel_row
            )
            start, end = _questionnaire_range(
                values[columns["questionnaire_range"]], excel_row
            )
            if end - start + 1 != target_n:
                raise PSUFrameError(
                    f"Baris {excel_row}: panjang rentang NO KUES ({end - start + 1}) "
                    f"tidak sama dengan RESPONDEN ({target_n})."
                )
            for previous_start, previous_end, previous_psu in intervals:
                if start <= previous_end and end >= previous_start:
                    raise PSUFrameError(
                        f"Baris {excel_row}: rentang NO KUES tumpang tindih dengan PSU {previous_psu}."
                    )
            intervals.append((start, end, psu_number))
            location_key = normalize_location(
                text_fields["village"], text_fields["district"], text_fields["regency"]
            )
            if location_key in seen_locations:
                warnings.append(
                    f"PSU {psu_number}: kombinasi desa, kecamatan, kabupaten/kota berulang."
                )
            seen_locations.add(location_key)
            parsed.append(
                {
                    "psu_number": psu_number,
                    **text_fields,
                    "target_n": target_n,
                    "questionnaire_start": start,
                    "questionnaire_end": end,
                    "normalized_location_key": location_key,
                }
            )
        if not parsed:
            raise PSUFrameError("Sheet PSU tidak memiliki baris data.")
        ordered_intervals = sorted(intervals)
        gaps = []
        for current, following in zip(ordered_intervals, ordered_intervals[1:]):
            if following[0] > current[1] + 1:
                gaps.append(f"{current[1] + 1}-{following[0] - 1}")
        if gaps:
            warnings.append("Nomor kuesioner memiliki jeda: " + ", ".join(gaps[:10]))
        target_total = sum(row["target_n"] for row in parsed)
        report = {
            "sheet": "PSU",
            "header_row": header_row,
            "row_count": len(parsed),
            "target_total": target_total,
            "questionnaire_start": min(item[0] for item in intervals),
            "questionnaire_end": max(item[1] for item in intervals),
            "regency_count": len({row["regency"] for row in parsed}),
            "dpr_ri_count": len({row["dpr_ri_constituency"] for row in parsed}),
            "warnings": warnings,
            "errors": [],
        }
        return PSUFrameCandidate(
            source_name=source_name,
            file_sha256=hashlib.sha256(raw).hexdigest(),
            rows=tuple(parsed),
            report=report,
        )
    finally:
        workbook.close()


def stage_psu_frame(*, survey, version, uploaded_file, user=None) -> SurveyPSUFrame:
    candidate = inspect_psu_frame(uploaded_file)
    with transaction.atomic():
        locked = SurveyAccess.objects.select_for_update().get(pk=survey.pk)
        if locked.psu_frames.filter(version=version).exists():
            raise PSUFrameError("Versi frame PSU sudah digunakan pada event ini.")
        if locked.psu_frames.filter(file_sha256=candidate.file_sha256).exists():
            raise PSUFrameError("File frame yang sama sudah pernah diunggah pada event ini.")
        try:
            frame = SurveyPSUFrame.objects.create(
                survey=locked,
                version=version,
                source_name=candidate.source_name,
                file_sha256=candidate.file_sha256,
                row_count=len(candidate.rows),
                target_total=candidate.report["target_total"],
                import_report=candidate.report,
                created_by=user,
            )
            SurveyPSU.objects.bulk_create(
                [SurveyPSU(frame=frame, **row) for row in candidate.rows]
            )
        except IntegrityError as exc:
            raise PSUFrameError("Frame PSU gagal disimpan karena konflik data.") from exc
    return frame


def activate_psu_frame(*, survey, frame_id) -> SurveyPSUFrame:
    with transaction.atomic():
        locked = SurveyAccess.objects.select_for_update().get(pk=survey.pk)
        frame = SurveyPSUFrame.objects.select_for_update().get(
            pk=frame_id, survey=locked
        )
        actual_rows = frame.psus.count()
        actual_target = sum(frame.psus.values_list("target_n", flat=True))
        if actual_rows != frame.row_count or actual_target != frame.target_total:
            raise PSUFrameError("Integritas frame berubah; aktivasi dibatalkan.")
        locked.psu_frames.filter(is_active=True).exclude(pk=frame.pk).update(is_active=False)
        frame.is_active = True
        frame.save(update_fields=("is_active",))
        source = locked.data_source
        if source.target_n != frame.target_total:
            source.target_n = frame.target_total
            source.save(update_fields=("target_n", "updated_at"))
    return frame
